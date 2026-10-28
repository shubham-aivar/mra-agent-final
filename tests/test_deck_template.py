"""The built-in default deck template, checking any deck against it, and making a clean template from a filled-in deck."""
import io
import os
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from mra.outputs import deck_template as dt
from mra.outputs import render
from mra.outputs.pptx_out import _lookup, render_pptx
from tests.test_outputs import model, template_bytes  # noqa: F401  (a small deck with some of the AWS deck's kinds of shapes)

DEFAULT = dt.default_template_bytes()
EXPECTED = {"cover_slides": 2, "charts": 20, "heatmap_boxes": 34, "focus_boxes": 6, "exec_boxes": 10, "activity_boxes": 6,
            "key_observation_lists": 2, "observation_rows": 38, "qa_rows": 84}           # the AWS deck: 62 scored + 22 scope/activity questions


def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def tables(prs):
    return [sh.table for s in prs.slides for sh in walk(s.shapes) if getattr(sh, "has_table", False) and sh.has_table]


def charts(prs):
    return [sh.chart for s in prs.slides for sh in walk(s.shapes) if getattr(sh, "has_chart", False) and sh.has_chart]


def all_text(prs):
    out = []
    for s in prs.slides:
        for sh in walk(s.shapes):
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                out += [c.text for r in sh.table.rows for c in r.cells]
    return "\n".join(out)


# ---- the default ------------------------------------------------------------------------------------------------------------
def test_the_default_deck_has_every_part_the_agent_fills(cfg):
    cov = dt.coverage(DEFAULT, cfg)
    assert cov["complete"] and cov["warnings"] == [] and cov["found"] == cov["expected"]
    assert {i["key"]: i["found"] for i in cov["items"]} == EXPECTED
    assert {i["key"]: i["expected"] for i in cov["items"]} == EXPECTED


def test_the_default_contains_no_answers_ratings_scores_or_observations():
    prs = Presentation(io.BytesIO(DEFAULT))
    qa = [t for t in tables(prs) if [c.text.strip() for c in t.rows[0].cells] == ["#", "Question", "Response", "Comments"]]
    assert len(qa) >= 10 and sum(len(t.rows) - 1 for t in qa) == EXPECTED["qa_rows"]
    assert all(r.cells[2].text == "" and r.cells[3].text == "" for t in qa for r in list(t.rows)[1:]), "a response or comment survived"
    obs = [t for t in tables(prs) if [c.text.strip() for c in t.rows[0].cells] == ["Category", "Summary Observations", "Actions & Next Steps"]]
    assert obs and all(r.cells[1].text == "" and r.cells[2].text == "" for t in obs for r in list(t.rows)[1:])
    for ch in charts(prs):
        for series in ch.plots[0].series:
            if series.name != "Guide":
                assert all(v is None for v in series.values), f"chart '{series.name}' still holds values"
            else:
                assert list(series.values) == [80.0] * len(list(series.values))             # the guide ring is the AWS standard, kept


def test_the_default_shows_placeholders_where_the_customer_goes():
    text = all_text(Presentation(io.BytesIO(DEFAULT)))
    for ph in ("[Customer Name]", "[Company Name]", "[Date]", "[Presenter]", "[Answer to Q1]", "[Answer to Q22]", "[Generated from the scores]"):
        assert ph in text, ph
    assert "FACILITATOR" in text                                                              # the AWS instruction banners stay; filling removes them


def test_the_default_carries_no_personal_names_comments_or_thumbnail():
    assert dt.personal_names(DEFAULT) == []
    with zipfile.ZipFile(io.BytesIO(DEFAULT)) as z:
        names = z.namelist()
        assert not any("thumbnail" in n or n.startswith("ppt/comments") for n in names)
        assert b"<p:cmAuthor " not in z.read("ppt/commentAuthors.xml")
        core = z.read("docProps/core.xml").decode()
        assert "<cp:lastModifiedBy></cp:lastModifiedBy>" in core and "<dc:creator></dc:creator>" in core


def test_the_default_is_a_valid_deck_python_can_reopen_and_save():
    prs = Presentation(io.BytesIO(DEFAULT))
    assert len(prs.slides) == 90
    buf = io.BytesIO()
    prs.save(buf)
    assert len(Presentation(io.BytesIO(buf.getvalue())).slides) == 90


# ---- filling the default -----------------------------------------------------------------------------------------------------
def test_filling_the_default_replaces_every_placeholder_and_follows_the_scores(svc, cfg, rset, full_ratings):
    m = model(svc, cfg, rset, full_ratings, answers={"Q1": "In scope text", "Q2": "Outcomes", "Q17": "Cloud first", "Q22": "No cloud team"})
    out = render("pptx", m, DEFAULT)                                                           # the AWS format, given explicitly
    prs = Presentation(io.BytesIO(out))
    text = all_text(prs)
    for leftover in ("[Customer Name]", "[Company Name]", "[Date]", "[Presenter]", "[Answer to", "[Generated from the scores]", "[Strength]",
                     "[Gap]", "FACILITATOR", "[IMAGE"):
        assert leftover not in text, leftover
    assert "Ingram Micro" in text and "Fri Oct 30 2026" in text and "A. Person" in text and "Cloud first" in text
    assert len(prs.slides) == 89 and m["deck_stats"]["slides_dropped"] == 1 and m["deck_warnings"] == []
    assert {k: m["deck_stats"][k] for k in EXPECTED} == EXPECTED
    bars = next(c for c in charts(prs) if c.plots[0].series[0].name == "Average Readiness Scores")
    by_ws = {w["name"]: w["score"] for w in m["workstreams"]}
    labelled = [(_lookup(cat, list(by_ws)), v) for cat, v in zip(bars.plots[0].categories, bars.plots[0].series[0].values)]
    assert len(labelled) == 8 and all(by_ws[name] == v for name, v in labelled)                # matched by name: the AWS bar chart lists them bottom-up
    radar = next(c for c in charts(prs) if c.plots[0].series[0].name == "Perspectives")
    by_p = {p["name"]: p["radar"] for p in m["perspectives"]}
    assert all(by_p[_lookup(cat, list(by_p))] == v for cat, v in zip(radar.plots[0].categories, radar.plots[0].series[0].values))
    qa = {r.cells[0].text: r for t in tables(prs) for r in list(t.rows)[1:] if r.cells[0].text.startswith("Q")}
    assert qa["Q26"].cells[2].text == "4" and qa["Q26"].cells[3].text.startswith("4 - ") and qa["Q22"].cells[3].text == "No cloud team"
    bc1 = next(a for a in m["activities"] if a["id"] == "BC-1")
    obs = next(r for t in tables(prs) for r in list(t.rows)[1:] if r.cells[0].text == bc1["name"] and len(r.cells) == 3)
    assert obs.cells[1].text == bc1["observation"]["text"]


def test_an_incomplete_group_stays_empty_in_a_filled_default(svc, cfg, rset, golden):
    m = model(svc, cfg, rset, golden["ratings"])                                                # Q42 is blank: People and Skills & COE are incomplete
    prs = Presentation(io.BytesIO(render("pptx", m, DEFAULT)))
    radar = next(c for c in charts(prs) if c.plots[0].series[0].name == "Perspectives")
    got = dict(zip(radar.plots[0].categories, radar.plots[0].series[0].values))
    assert got["People"] is None and all(v is not None for k, v in got.items() if k != "People")


def test_filling_never_changes_the_default_itself(svc, cfg, rset, full_ratings):
    before = dt.default_template_bytes()
    render("pptx", model(svc, cfg, rset, full_ratings), before)
    assert dt.default_template_bytes() == before and dt.coverage(before, cfg)["complete"]


# ---- checking another deck -----------------------------------------------------------------------------------------------------
def test_a_deck_with_only_some_parts_is_reported_as_partial(cfg, template_bytes):  # noqa: F811
    cov = dt.coverage(template_bytes, cfg)
    assert not cov["complete"] and 0 < cov["found"] < cov["expected"] and cov["warnings"]
    items = {i["key"]: i for i in cov["items"]}
    assert items["charts"]["found"] == 2 and items["charts"]["expected"] == 20 and not items["charts"]["ok"]
    assert items["heatmap_boxes"]["found"] == 0 and "found 0 of 34" in " ".join(cov["warnings"])
    assert "partly matches" in dt.describe(cov)


def test_a_deck_with_nothing_recognisable_says_so(cfg):
    buf = io.BytesIO()
    Presentation().save(buf)
    cov = dt.coverage(buf.getvalue(), cfg)
    assert cov["found"] == 0 and cov["warnings"] == ["this does not look like the AWS MRA readout deck: nothing in it can be filled"]


def test_a_file_that_is_not_a_deck_is_refused(cfg):
    with pytest.raises(ValueError, match="cannot be filled as a deck"):
        dt.coverage(b"not a deck", cfg)


def test_a_partial_deck_still_fills_what_it_has_and_the_open_items_say_what_it_lacks(svc, cfg, rset, full_ratings, template_bytes):  # noqa: F811
    svc.run(rset(full_ratings, meta={"facilitator": "F", "workshop_date": "2026-10-23", "readout_date": "2026-10-30"}), "t")
    files = svc.export("t", ["pptx", "todo"], "t", {"pptx": template_bytes}, {"pptx": "my deck.pptx"})
    todo = next(v for k, v in files.items() if k.endswith("open_items.md")).decode()
    assert "Deck template used: my deck.pptx" in todo and "Heatmap boxes: found 0 of 34" in todo
    last = [e for e in svc.audit_entries("t") if e["action"] == "export" and e["detail"]["file"].endswith(".pptx")][-1]["detail"]["template"]
    assert last["name"] == "my deck.pptx" and len(last["sha256"]) == 64


def test_the_audit_log_names_the_built_in_template(svc, cfg, rset, full_ratings):
    svc.run(rset(full_ratings), "t")
    svc.export("t", ["pptx"], "t")
    d = [e for e in svc.audit_entries("t") if e["action"] == "export" and e["detail"]["file"].endswith(".pptx")][0]["detail"]["template"]
    assert d == {"name": "Ingram MAP deliverable (built in)", "sha256": dt.template_hash(None)}


# ---- making a clean template ---------------------------------------------------------------------------------------------------
def test_template_mode_keeps_the_aws_banners_and_a_normal_fill_removes_them(cfg, template_bytes):  # noqa: F811
    out = render_pptx(dt.blank_model(cfg), template_bytes)
    assert "FACILITATOR" in all_text(Presentation(io.BytesIO(out))) and len(Presentation(io.BytesIO(out)).slides) == 6
    m = dt._complete_model(cfg)
    assert "FACILITATOR" not in all_text(Presentation(io.BytesIO(render_pptx(m, template_bytes))))


def test_making_a_template_removes_names_and_checks_the_package(cfg, template_bytes):  # noqa: F811
    prs = Presentation(io.BytesIO(template_bytes))
    prs.core_properties.last_modified_by = "Jane Doe"
    prs.core_properties.author = "John Roe"
    buf = io.BytesIO()
    prs.save(buf)
    src = _with_authors(buf.getvalue())                                                         # plus a comment-authors part, as the AWS deck has
    assert dt.personal_names(src) == ["Jane Doe", "John Roe", "Sam Poe"]
    out, report = dt.make_default_template(src, cfg)
    assert dt.personal_names(out) == [] and report["personal_names_removed"] == 3 and report["stats"]["charts"] == 2
    text = dt.package_text(out)
    assert not any(n in text for n in ("jane doe", "john roe", "sam poe"))


def _with_authors(src: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("ppt/commentAuthors.xml", '<p:cmAuthorLst xmlns:p="x"><p:cmAuthor id="0" name="Sam Poe"/></p:cmAuthorLst>')
    return buf.getvalue()


def test_the_leak_check_stops_the_build_when_forbidden_text_remains(cfg, template_bytes):  # noqa: F811
    with pytest.raises(ValueError, match="still contain"):
        dt.make_default_template(template_bytes, cfg, forbid=["MRA Readiness Scores"])         # a slide title the template keeps
    out, report = dt.make_default_template(template_bytes, cfg, forbid=["a customer answer that is not in the deck"])
    assert report["checked_forbidden"] == 1 and out


def test_aws_wording_is_not_mistaken_for_a_leak(cfg, template_bytes):  # noqa: F811
    q = next(q for q in cfg["questions"] if q.id == "Q2")
    phrase = "primary drivers"
    assert phrase in q.text.lower()
    out, report = dt.make_default_template(template_bytes, cfg, forbid=[phrase])               # a customer answer that reused AWS's words
    assert out and report["checked_forbidden"] == 1


def test_the_package_scan_reaches_notes_chart_data_workbooks_and_properties(template_bytes):  # noqa: F811
    prs = Presentation(io.BytesIO(template_bytes))
    prs.slides[2].notes_slide.notes_text_frame.text = "a hidden speaker note"
    prs.core_properties.comments = "a property comment"
    buf = io.BytesIO()
    prs.save(buf)
    text = dt.package_text(buf.getvalue())
    assert "a hidden speaker note" in text and "a property comment" in text
    assert "average readiness scores" in text                                                    # chart series name and embedded workbook
    assert dt.residue(buf.getvalue(), ["A HIDDEN SPEAKER NOTE", "not there at all"]) == ["A HIDDEN SPEAKER NOTE"]


@pytest.mark.skipif(not os.environ.get("MRA_TEMPLATE_DECK"), reason="set MRA_TEMPLATE_DECK to the reference AWS deck to rebuild the default")
def test_the_default_can_be_rebuilt_from_the_reference_deck_with_the_same_structure(cfg):
    src = Path(os.environ["MRA_TEMPLATE_DECK"]).read_bytes()
    out, report = dt.make_default_template(src, cfg)
    assert {k: report["stats"][k] for k in EXPECTED} == EXPECTED
    assert dt.coverage(out, cfg)["complete"] and dt.personal_names(out) == []
    assert dt.deck_lines(out) == dt.deck_lines(DEFAULT)                                          # the same visible text as the committed default
