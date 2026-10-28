"""The Excel, Word and PowerPoint outputs and the charts."""
import io
import os
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches

from mra.outputs import RENDERERS, filename, render
from mra.outputs.charts import bars_png, radar_png
from mra.report_model import build_report_model
from mra.review import workflow as wf
from mra.traceability import safe_cell

SID = "t"
META = {"facilitator": "Ingram Micro", "workshop_date": "2026-10-23", "readout_date": "2026-10-30", "presenter": "A. Person"}


def model(svc, cfg, rset, ratings, **kw):
    svc.run(rset(ratings, meta=META, **kw), "tester")
    res, state = svc.result(SID), wf.new_state(SID)
    return build_report_model(res, state, "2026-10-26T00:00:00+00:00", [], cfg)


@pytest.fixture
def m_blank(svc, cfg, rset, golden):
    return model(svc, cfg, rset, golden["ratings"], answers={"Q1": "In scope text", "Q2": "Outcomes", "Q7": "Yes", "Q13": "No"})


@pytest.fixture
def m_full(svc, cfg, rset, full_ratings):
    return model(svc, cfg, rset, full_ratings, answers={"Q1": "In scope text", "Q2": "Outcomes", "Q17": "Cloud first", "Q22": "No cloud team"})


# ---- registry ---------------------------------------------------------------
def test_file_names_and_unknown_kinds():
    assert filename("x", "DRAFT", "xlsx") == "MRA_x_DRAFT.xlsx" and filename("x", "FINAL", "todo") == "MRA_x_FINAL_open_items.md"
    assert set(RENDERERS) == {"xlsx", "docx", "pptx", "todo", "context"}
    with pytest.raises(ValueError, match="unknown export"):
        render("pdf", {})


def test_with_no_template_the_deck_is_the_ingram_deliverable(m_full):
    out = render("pptx", m_full)
    assert out[:2] == b"PK" and m_full["deck_stats"]["slides"] == 25 and m_full["deck_stats"]["mra_charts"] == 12


# ---- Excel ------------------------------------------------------------------
def test_the_first_three_sheets_have_the_aws_names_columns_and_order(m_blank, cfg):
    wb = load_workbook(io.BytesIO(render("xlsx", m_blank)))
    assert wb.sheetnames[:3] == ["Questions", "Actions", "Scores"] and "Summary" in wb.sheetnames
    assert [c.value for c in wb["Questions"][1]] == ["Section", "Category", "#", "Question", "Response Type", "Response",
                                                    "Rating Guide (if applicable)", "Comments", "Rating 1", "Rating 2", "Rating 3", "Rating 4", "Rating 5"]
    assert [c.value for c in wb["Actions"][1]] == ["Section", "Readiness Activity", "Workstream Phase", "Summary Observations", "Action",
                                                  "Priority", "Effort", "Target Completion Date", "Benefit", "Owner", "Status"]
    assert [c.value for c in wb["Scores"][1]] == ["Section", "Category", "Score"]
    assert wb["Questions"].max_row == 163 and wb["Scores"].max_row == 73


def test_answers_land_where_the_aws_export_puts_them(m_blank):
    ws = load_workbook(io.BytesIO(render("xlsx", m_blank)))["Questions"]
    rows = {r[2].value: r for r in ws.iter_rows(min_row=2)}
    assert rows[1][7].value == "In scope text" and rows[1][5].value is None          # text answer in Comments
    assert rows[7][5].value == "Yes" and rows[13][5].value == "No"                   # Yes/No in Response
    assert rows[26][5].value == 4 and rows[26][6].value.startswith("4 - ")           # rating and its guide text
    assert rows[42][5].value is None and rows[42][6].value is None                   # the blank stays blank


def test_a_blank_rating_gives_empty_score_cells_never_zero(m_blank):
    sc = load_workbook(io.BytesIO(render("xlsx", m_blank)))["Scores"]
    rows = {(r[0].value, r[1].value): r[2].value for r in sc.iter_rows(min_row=2)}
    assert rows[("People", None)] is None and rows[("People", "Organizational Structure")] is None
    assert rows[("Business", None)] == 3.0 and rows[("Operations", None)] == 3.4


def test_actions_sheet_has_one_row_per_action_with_the_phase(m_full):
    ws = load_workbook(io.BytesIO(render("xlsx", m_full)))["Actions"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == len(m_full["action_rows"]) and rows[0][2] in ("Pre-Mobilize", "Mobilize") and rows[0][10] == "Proposed"


def test_text_that_starts_like_a_formula_is_never_a_formula(svc, cfg, rset, golden):
    """F64: a respondent's text such as =HYPERLINK(...) used to become a live formula."""
    m = model(svc, cfg, rset, golden["ratings"], answers={"Q1": '=HYPERLINK("http://x","y")', "Q2": "+cmd", "Q3": "@SUM(1)"})
    wb = load_workbook(io.BytesIO(render("xlsx", m)))
    live = [c.coordinate for ws in wb for row in ws.iter_rows() for c in row if c.data_type == "f" and ws.title not in ("Summary",)]
    assert live == []
    assert wb["Questions"].cell(2, 8).value == "'" + '=HYPERLINK("http://x","y")'
    assert safe_cell("=1+1") == "'=1+1" and safe_cell("-5") == "'-5" and safe_cell("ok") == "ok" and safe_cell(4) == 4


def test_the_summary_formulas_count_the_other_sheets(m_full):
    ws = load_workbook(io.BytesIO(render("xlsx", m_full)))["Summary"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert len(formulas) == 5 and any("Traceability" in f for f in formulas) and any("To validate" in f for f in formulas)


def test_adjustments_get_their_own_sheet(svc, cfg, rset, golden):
    svc.run(rset(golden["ratings"], meta=META), "tester")
    svc.adjust(SID, "Q42", 1, "from the notes", "tc")
    m = build_report_model(svc.result(SID), wf.new_state(SID), "t", [], cfg)
    ws = load_workbook(io.BytesIO(render("xlsx", m)))["Adjustments"]
    assert [c.value for c in ws[2]][:4] == ["Q42", None, 1, "from the notes"]


# ---- Word -------------------------------------------------------------------
def _text(doc):
    return "\n".join(p.text for p in doc.paragraphs)


def test_the_word_report_has_the_aws_sections_in_order(m_full):
    doc = Document(io.BytesIO(render("docx", m_full)))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading 1")]
    assert heads == ["Overview", "Contents", "Radar charts", "Heatmap", "Scores", "Details", "Consolidated recommended actions", "Responses",
                     "Method and limits"]
    assert len(doc.inline_shapes) == 2 + 8                                             # two radars and eight workstream bar charts
    t = _text(doc)
    assert "Summary Observations:" in t and "Recommended Actions:" in t and "Facilitated by: Ingram Micro" in t


def test_incomplete_is_stated_never_zero(m_blank):
    doc = Document(io.BytesIO(render("docx", m_blank)))
    t = _text(doc) + "\n".join(c.text for tb in doc.tables for r in tb.rows for c in r.cells)
    assert "Incomplete: 1 rating(s) are blank (Q42)" in t and "Incomplete" in t
    scores = [[c.text for c in r.cells] for r in doc.tables[1].rows]
    assert ["People", "Incomplete", "5", "6"] in scores


def test_missing_report_details_print_tbd_not_a_template_placeholder(svc, cfg, rset, golden):
    svc.run(rset(golden["ratings"]), "tester")
    m = build_report_model(svc.result(SID), wf.new_state(SID), "t", [], cfg)
    t = _text(Document(io.BytesIO(render("docx", m))))
    assert "Facilitated by: TBD" in t and "[" not in t.split("Overview")[0]


def test_every_response_is_listed(m_full):
    doc = Document(io.BytesIO(render("docx", m_full)))
    cells = "\n".join(c.text for tb in doc.tables for r in tb.rows for c in r.cells)
    assert "Cloud first" in cells and "No cloud team" in cells and cells.count("Do you have a business case") >= 1


def test_a_word_template_keeps_its_styles_and_header(m_full):
    base = Document()
    base.sections[0].header.paragraphs[0].text = "AWS HEADER"
    base.add_paragraph("old body text that must go")
    buf = io.BytesIO()
    base.save(buf)
    doc = Document(io.BytesIO(render("docx", m_full, buf.getvalue())))
    assert doc.sections[0].header.paragraphs[0].text == "AWS HEADER" and "old body text" not in _text(doc)


# ---- charts -----------------------------------------------------------------
def test_charts_are_png_images_and_handle_incomplete_values():
    from PIL import Image
    r = radar_png("Perspectives", ["A", "B", "C", "D", "E", "F"], [60, 70, None, 50, 80, 40])
    b = bars_png("Business Case", [("AVERAGE", 3.2, "amber"), ("Activity one", 4.0, "green"), ("Activity two", None, "incomplete")])
    for png, size in ((r, (640, 480)), (b, (700, 34 + 26 * 3 + 16))):
        im = Image.open(io.BytesIO(png))
        assert im.format == "PNG" and im.size == size
    assert radar_png("x", ["A", "B", "C"], [1, 2, 3]) == radar_png("x", ["A", "B", "C"], [1, 2, 3])


# ---- PowerPoint: a small template with the same kinds of shapes as the AWS deck -------------
def _add_title(slide, text):
    box = slide.shapes.add_textbox(Emu(500_000), Emu(100_000), Emu(12_000_000), Emu(600_000))
    box.text_frame.text = text
    return box


def _chart(slide, name, cats, vals, kind, top=1_000_000, extra_series=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(name, vals)
    for n, v in (extra_series or []):
        cd.add_series(n, v)
    return slide.shapes.add_chart(kind, Emu(300_000), Emu(top), Emu(5_000_000), Emu(2_500_000), cd).chart


@pytest.fixture
def template_bytes(cfg):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(9)
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)                                                      # guidance slide: removed
    _add_title(s, "Guidance for using this deck")
    s = prs.slides.add_slide(blank)                                                      # cover
    _add_title(s, "Migration Readiness Assessment (MRA)")
    for text, top in (("Old Customer", 2_000_000), ("[Company Name]", 5_500_000), ("Fri Jan 01 2021", 6_200_000), ("role/SSO_Admin_123", 6_800_000)):
        b = s.shapes.add_textbox(Emu(600_000), Emu(top), Emu(4_000_000), Emu(400_000))
        b.text_frame.text = text
    f = s.shapes.add_textbox(Emu(600_000), Emu(1_200_000), Emu(2_000_000), Emu(300_000))
    f.text_frame.text = "Facilitated By"
    s = prs.slides.add_slide(blank)                                                      # average scores bar chart
    _add_title(s, "MRA Readiness Scores")
    _chart(s, "Average Readiness Scores", [w["name"][:12] + "..." if len(w["name"]) > 12 else w["name"] for w in cfg["workstreams"]],
           [1.0] * 8, XL_CHART_TYPE.BAR_CLUSTERED)
    b = s.shapes.add_textbox(Emu(100_000), Emu(8_000_000), Emu(6_000_000), Emu(300_000))
    b.text_frame.text = "FACILITATOR: delete me"
    s = prs.slides.add_slide(blank)                                                      # perspective radar + three sentences
    _add_title(s, "MRA Areas of Focus")
    _chart(s, "Perspectives", [p["name"] for p in cfg["perspectives"]], [10] * 6, XL_CHART_TYPE.RADAR_FILLED, extra_series=[("Guide", [80] * 6)])
    for i in range(3):
        r = s.shapes.add_textbox(Emu(8_500_000), Emu(1_800_000 + i * 1_900_000), Emu(5_500_000), Emu(1_500_000))
        r.text_frame.text = f"old sentence {i}"
    s = prs.slides.add_slide(blank)                                                      # observation table
    _add_title(s, "Business Case")
    t = s.shapes.add_table(5, 3, Emu(500_000), Emu(1_400_000), Emu(12_000_000), Emu(2_000_000)).table
    for j, h in enumerate(["Category", "Summary Observations", "Actions & Next Steps"]):
        t.cell(0, j).text = h
    for i, name in enumerate(["High Level Business Case", "Key Stakeholder Sign-off", "Migration Funding Commitment",
                              "Specific Migration Workloads Committed"], 1):
        t.cell(i, 0).text = name
        t.cell(i, 1).text = "old observation"
    s = prs.slides.add_slide(blank)                                                      # question and answer table
    _add_title(s, "Business Case - High Level Business Case")
    t = s.shapes.add_table(4, 4, Emu(500_000), Emu(1_400_000), Emu(12_000_000), Emu(2_000_000)).table
    for j, h in enumerate(["#", "Question", "Response", "Comments"]):
        t.cell(0, j).text = h
    for i, q in enumerate(["Q26", "Q29", "Q33"], 1):
        t.cell(i, 0).text, t.cell(i, 2).text, t.cell(i, 3).text = q, "9", "old comment"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _slide_text(prs):
    return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)


def test_the_deck_is_filled_from_the_model(m_full, template_bytes):
    out = render("pptx", m_full, template_bytes)
    prs = Presentation(io.BytesIO(out))
    assert len(prs.slides) == 5                                   # the guidance slide is gone
    text = _slide_text(prs)
    assert "FACILITATOR" not in text and "old sentence" not in text and "[Company Name]" not in text and "SSO_Admin" not in text
    assert "Ingram Micro" in text and "Fri Oct 30 2026" in text and "A. Person" in text and "Customer A" in text
    st = m_full["deck_stats"]
    assert st["charts"] == 2 and st["observation_rows"] == 4 and st["qa_rows"] == 3 and st["slides_dropped"] == 1


def test_chart_values_and_embedded_workbooks_follow_the_scores(m_full, template_bytes):
    prs = Presentation(io.BytesIO(render("pptx", m_full, template_bytes)))
    charts = {sh.chart.plots[0].series[0].name: sh.chart for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart}
    bars = charts["Average Readiness Scores"]
    want = [w["score"] for w in m_full["workstreams"]]
    assert list(bars.plots[0].series[0].values) == want
    emb = load_workbook(io.BytesIO(bars.part.chart_workbook.xlsx_part.blob)).active
    assert [emb.cell(r, 2).value for r in range(2, 10)] == want
    radar = charts["Perspectives"]
    assert list(radar.plots[0].series[0].values) == [p["radar"] for p in m_full["perspectives"]] and list(radar.plots[0].series[1].values) == [80] * 6
    from pptx.dml.color import RGBColor
    for i, w in enumerate(m_full["workstreams"]):
        want_hex = {"green": "00CC00", "amber": "FFFF66", "red": "FF9999"}[w["colour"]]
        assert str(bars.plots[0].series[0].points[i].format.fill.fore_color.rgb) == str(RGBColor.from_string(want_hex))


def test_incomplete_groups_have_no_bar_and_no_radar_point(m_blank, template_bytes):
    prs = Presentation(io.BytesIO(render("pptx", m_blank, template_bytes)))
    charts = {sh.chart.plots[0].series[0].name: sh.chart for s in prs.slides for sh in s.shapes if getattr(sh, "has_chart", False) and sh.has_chart}
    names = [w["name"] for w in m_blank["workstreams"]]
    assert list(charts["Average Readiness Scores"].plots[0].series[0].values)[names.index("Skills & COE")] is None
    assert list(charts["Perspectives"].plots[0].series[0].values)[2] is None


def test_tables_get_the_observations_actions_and_answers(m_full, template_bytes):
    prs = Presentation(io.BytesIO(render("pptx", m_full, template_bytes)))
    tables = [sh.table for s in prs.slides for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table]
    obs, qa = tables
    bc1 = next(a for a in m_full["activities"] if a["id"] == "BC-1")
    assert obs.cell(1, 1).text == bc1["observation"]["text"] and "old observation" not in "".join(c.text for r in obs.rows for c in r.cells)
    assert obs.cell(1, 2).text and qa.cell(1, 2).text == "4" and qa.cell(1, 3).text.startswith("4 - ") and qa.cell(3, 2).text == "2"


def test_the_focus_sentences_replace_the_old_ones(m_full, template_bytes):
    text = _slide_text(Presentation(io.BytesIO(render("pptx", m_full, template_bytes))))
    for sent in m_full["focus"]["perspective_sentences"]:
        assert sent in text


def test_a_template_that_still_names_another_customer_is_reported(m_full, template_bytes):
    """A previously filled deck used as the template: the old customer's name must not survive unnoticed."""
    prs = Presentation(io.BytesIO(template_bytes))
    prs.slides[2].shapes.add_textbox(Emu(100_000), Emu(7_500_000), Emu(4_000_000), Emu(300_000)).text_frame.text = "Old Customer plan"
    buf = io.BytesIO()
    prs.save(buf)
    render("pptx", m_full, buf.getvalue())
    assert any("Old Customer" in w for w in m_full["deck_warnings"])
    m_full["deck_warnings"] = None
    render("pptx", m_full, template_bytes)
    assert m_full["deck_warnings"] == []


def test_the_open_items_file_lists_what_a_person_must_do(m_blank, m_full):
    t = render("todo", m_blank).decode()
    assert "Blank ratings" in t and "Q42" in t and "Mobilize cost slides" in t and "AMS fit score" in t and "DRAFT. DRAFT." not in t
    assert "Blank ratings" not in render("todo", m_full).decode()


def test_the_planning_note_lists_gaps_and_accepted_items(m_full):
    m_full["context_items"] = [{"kind": "constraint", "text": "Lease ends June 2027", "evidence_id": "E1", "quote": "the lease ends", "id": "P1"}]
    t = render("context", m_full).decode()
    assert "Lease ends June 2027" in t and "Source: E1" in t and "## Readiness gaps" in t


@pytest.mark.skipif(not os.environ.get("MRA_TEMPLATE_DECK"), reason="set MRA_TEMPLATE_DECK to the real AWS deck to run this")
def test_the_real_aws_deck_is_filled_without_leaving_the_old_customers_text(m_full):
    out = render("pptx", m_full, Path(os.environ["MRA_TEMPLATE_DECK"]).read_bytes())
    prs = Presentation(io.BytesIO(out))
    assert "FACILITATOR" not in _slide_text(prs) and m_full["deck_stats"]["charts"] == 20 and m_full["deck_stats"]["heatmap_boxes"] == 34
