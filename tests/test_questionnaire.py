"""Reading the filled-in questionnaire, and building the question bank from the template."""
import json

import pytest
from openpyxl import load_workbook

from mra.inputs.evidence import load_evidence
from mra.inputs.question_import import import_questions
from mra.inputs.questionnaire import load_input, read_questionnaire
from mra.loader import InputValidationError, build_response_set, read_json
from tests.conftest import FIXTURES, write_questionnaire


def test_a_filled_questionnaire_round_trips(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"], {1: "scope text", 2: "outcomes", 7: "Yes", 13: "No", 98: "yes"})
    rs, hashes = load_input(p, cfg, customer_label="Customer A")
    assert len(rs.ratings) == 61 and {r.question_id for r in rs.ratings}.isdisjoint({"Q42"})
    got = {a.question_id: a.value for a in rs.answers}
    assert got["Q1"] == "scope text" and got["Q7"] == "Yes" and got["Q13"] == "No" and got["Q98"] == "yes"
    assert rs.set_id == "q" and list(hashes) == ["q.xlsx"]


def test_a_blank_rating_stays_blank_never_zero(cfg, golden, tmp_path):
    rs, _ = load_input(write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"]), cfg)
    assert "Q42" not in {r.question_id for r in rs.ratings}


@pytest.mark.parametrize("bad", ["3.5", "4.9", "2,5", "5 out of 10", "3/5", "three", "0", "6", "-1"])
def test_a_rating_that_is_not_a_whole_number_from_1_to_5_is_an_error(cfg, golden, tmp_path, bad):
    """F63: '3.5' used to become 3 and '5 out of 10' became 5, silently."""
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"], overrides={26: {"Response": bad}})
    with pytest.raises(InputValidationError, match=r"row \d+: rating for #26"):
        read_questionnaire(p, cfg)


def test_a_whole_number_stored_as_a_float_is_fine(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"], overrides={26: {"Response": 4.0}})
    assert {r["question_id"]: r["score"] for r in read_questionnaire(p, cfg)["ratings"]}["Q26"] == 4


def test_the_rating_guide_level_is_read_and_a_disagreement_is_kept(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"], overrides={26: {"Response": 2}})
    rs, _ = load_input(p, cfg)
    r = next(x for x in rs.ratings if x.question_id == "Q26")
    assert (r.score, r.guide_level) == (2, 4)


def test_all_problems_are_reported_together(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"],
                            overrides={26: {"Response": 9}, 27: {"Response": "x"}, 7: {"Response": "maybe"}, 29: {"Response Type": "Text"}})
    with pytest.raises(InputValidationError) as e:
        read_questionnaire(p, cfg)
    assert len(e.value.errors) == 4


def test_unknown_question_numbers_and_duplicates_are_errors(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"])
    wb = load_workbook(p)
    ws = wb["Questions"]
    ws.append(["Business", "X", 999, "new question", "Rating", 3])
    ws.append(["Business", "X", 26, "again", "Rating", 3])
    wb.save(p)
    with pytest.raises(InputValidationError) as e:
        read_questionnaire(p, cfg)
    text = " ".join(e.value.errors)
    assert "#999 is not in the question bank" in text and "#26 appears twice" in text


def test_changed_wording_is_a_warning_not_an_error(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "q.xlsx", cfg, golden["ratings"])
    wb = load_workbook(p)
    wb["Questions"].cell(2 + 25, 4).value = "A reworded question"      # row of #26
    wb.save(p)
    assert any("#26" in w and "wording differs" in w for w in read_questionnaire(p, cfg)["warnings"])


def test_wrong_headers_say_what_is_missing(cfg, tmp_path):
    p = tmp_path / "q.csv"
    p.write_text("A,B\n1,2\n")
    with pytest.raises(InputValidationError, match="no column found"):
        read_questionnaire(p, cfg)


def test_the_canonical_json_input_and_meta_are_validated(cfg, golden):
    ok = {"set_id": "j", "customer_label": "J", "ratings": [{"question_id": "Q26", "score": 4}], "answers": [], "evidence": [],
          "meta": {"facilitator": "F", "readout_date": "2026-11-01"}}
    assert build_response_set(ok, cfg).meta == {"facilitator": "F", "readout_date": "2026-11-01"}
    bad = {**ok, "ratings": [{"question_id": "Q26", "score": 9}, {"question_id": "Q26", "score": 3}, {"question_id": "Q1", "score": 3}],
           "meta": {"nonsense": "x"}}
    with pytest.raises(InputValidationError) as e:
        build_response_set(bad, cfg)
    assert len(e.value.errors) == 4
    with pytest.raises(InputValidationError, match="set_id"):
        build_response_set({**ok, "set_id": "../x"}, cfg)


def test_evidence_can_be_tagged_to_a_question_a_category_or_an_activity(cfg):
    base = {"set_id": "e", "customer_label": "E", "ratings": [{"question_id": "Q26", "score": 4}], "answers": []}
    for target in ("Q26", "BUS-04", "BC-1"):
        build_response_set({**base, "evidence": [{"evidence_id": "E1", "text": "x", "applies_to": [target]}]}, cfg)
    with pytest.raises(InputValidationError, match="applies_to"):
        build_response_set({**base, "evidence": [{"evidence_id": "E1", "text": "x", "applies_to": ["Q999"]}]}, cfg)


def test_evidence_from_a_folder_a_table_and_json(cfg, tmp_path):
    d = tmp_path / "notes"
    d.mkdir()
    (d / "cio.md").write_text("applies_to: Q26, BC-2\nThe CIO approved the budget.")
    (d / "other.txt").write_text("plain note")
    ev = load_evidence(d, cfg)
    assert [e["evidence_id"] for e in ev] == ["E-cio", "E-other"] and ev[0]["applies_to"] == ["Q26", "BC-2"]
    (d / "big.txt").write_text("x" * 250_000)
    with pytest.raises(InputValidationError, match="larger than"):
        load_evidence(d, cfg)
    t = tmp_path / "ev.csv"
    t.write_text("Evidence ID,Source,Text,Applies to\nE9,call,hello,Q26;Q27\n")
    assert load_evidence(t, cfg)[0]["applies_to"] == ["Q26", "Q27"]
    j = tmp_path / "ev.json"
    j.write_text(json.dumps({"evidence": [{"evidence_id": "E1", "text": "t"}]}))
    assert load_evidence(j, cfg)[0]["evidence_id"] == "E1" and load_evidence(None, cfg) == []


# ---- the importer ----
def _template(cfg, tmp_path):
    return write_questionnaire(tmp_path / "template.xlsx", cfg, {})


def test_the_importer_rebuilds_the_bank_from_the_template_exactly(cfg, tmp_path):
    bank, rep = import_questions(_template(cfg, tmp_path), cfg["mapping"], "t")
    saved = read_json(FIXTURES / "config" / "question_bank.json")
    assert bank["questions"] == saved["questions"] and bank["areas"] == saved["areas"] and bank["modules"] == saved["modules"]
    assert rep["questions"] == 162 and rep["scored_ratings"] == 62 and rep["by_type"] == {"rating": 68, "yes_no": 19, "text": 75}


def test_the_importer_ignores_answers(cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "filled.xlsx", cfg, golden["ratings"], {1: "SECRET customer text"})
    assert "SECRET" not in json.dumps(import_questions(p, cfg["mapping"], "t")[0])


def test_the_importer_reports_every_problem(cfg, tmp_path):
    p = write_questionnaire(tmp_path / "t.xlsx", cfg, {})
    wb = load_workbook(p)
    ws = wb["Questions"]
    ws.cell(2, 1).value = "Mystery"
    ws.cell(3, 3).value = ws.cell(4, 3).value
    ws.cell(4 + 22, 9).value = None
    wb.save(p)
    with pytest.raises(InputValidationError) as e:
        import_questions(p, cfg["mapping"], "t")
    text = " ".join(e.value.errors)
    assert "unknown section 'Mystery'" in text and "duplicate question number" in text and "five rating-level descriptions" in text
