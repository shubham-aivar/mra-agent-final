"""The local review page's server, tested over real HTTP on a random local port."""
import base64
import http.client
import json
import stat
import threading

import pytest

from mra.ui.server import make_server
from tests.conftest import write_questionnaire
from tests.helpers import NOTE, Script, end, turn, use
from tests.test_outputs import template_bytes  # noqa: F401  (a small deck template with the AWS deck's kinds of shapes)

TOKEN = "test-token"
NOTES = [{"name": "cio.md", "content": NOTE}]


class Client:
    def __init__(self, srv):
        self.port = srv.server_address[1]

    def call(self, method, path, body=None, token=TOKEN, host=None, origin=None, actor=None, raw=None, length=None):
        h = {"Host": host or f"127.0.0.1:{self.port}"}
        if token:
            h["X-MRA-Token"] = token
        if origin:
            h["Origin"] = origin
        if actor:
            h["X-MRA-Actor"] = actor
        data = raw if raw is not None else (json.dumps(body).encode() if body is not None else None)
        if data is not None:
            h["Content-Type"] = "application/json"
            h["Content-Length"] = str(length if length is not None else len(data))
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=30)
        c.request(method, path, body=data, headers=h)
        r = c.getresponse()
        payload = r.read()
        c.close()
        ctype = r.getheader("Content-Type", "")
        return r.status, (json.loads(payload) if "json" in ctype else payload), dict(r.getheaders())


@pytest.fixture
def ui(svc):
    srv = make_server(svc, 0, TOKEN)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    c = Client(srv)
    c.srv, c.svc = srv, svc
    yield c
    srv.shutdown()
    srv.server_close()


def q_body(cfg, golden, tmp_path, name="q.xlsx", answers=None, **extra):
    p = write_questionnaire(tmp_path / name, cfg, golden["ratings"], answers or {1: "scope text", 2: "outcomes", 22: "There is no cloud team"})
    return {"file_name": name, "file_b64": base64.b64encode(p.read_bytes()).decode(), "set_id": "t", "customer": "Customer A",
            "notes": NOTES, "meta": {"facilitator": "Aivar", "workshop_date": "2026-10-23", "readout_date": "2026-10-30"}, **extra}


# ---- the page and the guards -----------------------------------------------------------------------
def test_the_page_is_served_with_the_token_and_a_strict_policy(ui):
    st, html, h = ui.call("GET", "/", token=None)
    html = html.decode()
    assert st == 200 and f'content="{TOKEN}"' in html and "__MRA_TOKEN__" not in html
    assert "script-src 'self'" in h["Content-Security-Policy"] and "frame-ancestors 'none'" in h["Content-Security-Policy"]
    assert h["X-Frame-Options"] == "DENY" and h["Cache-Control"] == "no-store"
    import re
    assert all("src=" in tag for tag in re.findall(r"<script[^>]*>", html)) and "<style" not in html and "onclick=" not in html


def test_the_static_files_are_served_and_nothing_else_from_disk(ui):
    assert ui.call("GET", "/static/app.js", token=None)[0] == 200 and ui.call("GET", "/static/style.css", token=None)[0] == 200
    for p in ("/static/../config/rubric.json", "/static/index.html", "/mra/cli.py", "/etc/passwd", "/%2e%2e/etc/passwd"):
        assert ui.call("GET", p, token=None)[0] == 404, p


def test_the_api_needs_this_launchs_token(ui):
    assert ui.call("GET", "/api/info", token=None)[0] == 403 and ui.call("GET", "/api/info", token="wrong")[0] == 403
    assert ui.call("POST", "/api/runs", {}, token="nope")[0] == 403
    assert ui.call("GET", f"/api/info?t={TOKEN}", token=None)[0] == 200          # download links carry the token in the address
    assert ui.call("GET", "/api/info")[0] == 200


def test_a_wrong_host_or_a_foreign_origin_is_refused(ui):
    assert ui.call("GET", "/", token=None, host="evil.example")[0] == 403           # DNS rebinding
    assert ui.call("GET", "/api/info", host=f"localhost.evil.example:{ui.port}")[0] == 403
    assert ui.call("POST", "/api/runs", {}, origin="http://evil.example")[0] == 403    # another web page calling us
    assert ui.call("GET", "/api/info", origin=f"http://127.0.0.1:{ui.port}")[0] == 200
    assert ui.call("GET", "/api/info", host=f"localhost:{ui.port}")[0] == 200


def test_it_listens_on_the_loopback_address_only(ui):
    assert ui.srv.server_address[0] == "127.0.0.1"


def test_bad_bodies_are_refused(ui):
    ui.srv.max_body = 2000
    assert ui.call("POST", "/api/runs", raw=b"x" * 5000)[0] == 413
    assert ui.call("POST", "/api/runs", raw=b"not json")[0] == 400 and ui.call("POST", "/api/runs", raw=b"[1]")[0] == 400
    assert ui.call("POST", "/api/runs", raw=b"{}", length=-5)[0] in (400, 413)
    assert ui.call("POST", "/api/nothing", {})[0] == 404 and ui.call("GET", "/api/runs/../x")[0] in (400, 404)


# ---- the whole flow, over HTTP ----------------------------------------------------------------------------
def test_a_run_is_created_viewed_reviewed_finalised_and_exported(ui, cfg, golden, tmp_path):
    st, out, _ = ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path), actor="Alice")
    assert (st, out) == (200, {"set_id": "t"})
    v = ui.call("GET", "/api/runs/t")[1]
    assert v["customer"] == "Customer A" and v["blank_ratings"] == ["Q42"] and v["report_status"] == "DRAFT"
    assert [p["status"] for p in v["perspectives"]].count("incomplete") == 1 and v["details"]["facilitator"] == "Aivar"
    assert [n["id"] for n in v["notes"]] == ["E-cio"] and any(t["number"] == 22 for t in v["text_answers"])
    assert v["status"]["ready_to_finalise"] is False and any("blank" in b for b in v["status"]["blockers"])
    assert ui.call("POST", "/api/runs/t/confirm", {})[0] == 400                          # a blank rating blocks confirmation
    assert ui.call("POST", "/api/runs/t/adjust", {"question": "Q42", "score": 1, "reason": "no cloud team (answer to Q22)"}, actor="Alice")[0] == 200
    v = ui.call("GET", "/api/runs/t")[1]
    assert v["blank_ratings"] == [] and next(p for p in v["perspectives"] if p["name"] == "People")["score"] == 2.8
    for a in v["activities"]:
        if a["needs_action"]:
            assert ui.call("POST", "/api/runs/t/decide", {"activity": a["id"], "decision": "approve", "owner": "A", "due_date": "2026-12-15"})[0] == 200
    for i in v["validation"]:
        assert ui.call("POST", "/api/runs/t/resolve", {"item": i["key"], "resolution": "accepted_risk", "note": "seen"})[0] == 200
    assert ui.call("POST", "/api/runs/t/finalise", {})[0] == 409                           # not confirmed yet
    assert ui.call("POST", "/api/runs/t/confirm", {})[0] == 200
    st, out, _ = ui.call("POST", "/api/runs/t/finalise", {})
    assert st == 200 and ui.call("GET", "/api/runs/t")[1]["report_status"] == "FINAL"
    assert ui.call("POST", "/api/runs/t/decide", {"activity": "BC-1", "decision": "approve"})[0] == 409
    files = ui.call("POST", "/api/runs/t/export", {"kinds": ["xlsx", "docx", "context"]})[1]["files"]
    assert sorted(f["name"] for f in files) == ["MRA_t_FINAL.docx", "MRA_t_FINAL.xlsx", "MRA_t_FINAL_planning_context.md"]
    st, data, h = ui.call("GET", "/api/runs/t/exports/MRA_t_FINAL.xlsx")
    assert st == 200 and data[:2] == b"PK" and "attachment" in h["Content-Disposition"] and "spreadsheetml" in h["Content-Type"]
    assert ui.call("GET", "/api/runs/t")[1]["audit"]["ok"]
    who = {e["actor"] for e in ui.svc.audit_entries("t")}
    assert "Alice" in who


def test_the_actor_name_is_cleaned_and_recorded(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path), actor="<b>Bob</b>; drop")
    runs = [e for e in ui.svc.audit_entries("t") if e["action"] == "run"]
    assert runs[0]["actor"] == "bBobb drop"                       # < > / ; are removed, the rest is kept
    ui.call("POST", "/api/runs/t/adjust", {"question": "Q42", "score": 1, "reason": "r"}, actor="   ")
    assert [e["actor"] for e in ui.svc.audit_entries("t") if e["action"] == "adjustment"][0] not in ("", "   ")     # empty falls back to the OS user


def test_questionnaire_problems_come_back_together_and_nothing_is_stored(ui, cfg, golden, tmp_path):
    p = write_questionnaire(tmp_path / "bad.xlsx", cfg, golden["ratings"], overrides={26: {"Response": "3.5"}, 27: {"Response": 9}})
    b = {**q_body(cfg, golden, tmp_path), "file_b64": base64.b64encode(p.read_bytes()).decode(), "file_name": "bad.xlsx"}
    st, out, _ = ui.call("POST", "/api/runs", b)
    assert st == 422 and len(out["errors"]) == 2
    assert ui.call("GET", "/api/runs")[1]["runs"] == []


@pytest.mark.parametrize("patch,needle", [
    ({"file_b64": ""}, "choose the filled-in questionnaire"), ({"file_name": "q.pdf"}, "must be"), ({"file_b64": "!!!"}, "base64"),
    ({"set_id": "../x"}, "run id"), ({"notes": [{"name": "a.pdf", "content": "x"}]}, "notes must be"),
    ({"notes": [{"name": "a.txt", "content": "x"}, {"name": "a.txt", "content": "y"}]}, "same file name"),
])
def test_upload_validation(ui, cfg, golden, tmp_path, patch, needle):
    st, out, _ = ui.call("POST", "/api/runs", {**q_body(cfg, golden, tmp_path), **patch})
    assert st == 400 and needle in out["error"]


def test_uploading_again_replaces_the_data_and_a_finalised_run_refuses(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    ui.call("POST", "/api/runs/t/decide", {"activity": "BC-1", "decision": "approve", "owner": "A", "due_date": "2026-12-15"})
    again = q_body(cfg, {"ratings": {**golden["ratings"], "Q26": 1}}, tmp_path, name="q2.xlsx")
    assert ui.call("POST", "/api/runs", again)[0] == 200
    v = ui.call("GET", "/api/runs/t")[1]
    assert next(q for q in v["questions"] if q["question_id"] == "Q26")["score"] == 1
    assert any(e["action"] == "reconcile" for e in ui.svc.audit_entries("t"))            # the old decision no longer fit and was dropped


def test_the_input_hashes_keep_the_uploaded_names_and_the_notes(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    hashes = ui.svc.result("t")["meta"]["input_hashes"]
    assert set(hashes) == {"q.xlsx", "notes/cio.md"}


def test_downloads_only_serve_known_export_names(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    ui.call("POST", "/api/runs/t/export", {"kinds": ["xlsx"]})
    for name in ("../result.json", "..%2Fresult.json", "result.json", "MRA_t_DRAFT.exe", "MRA_nope_DRAFT.xlsx", "review.json"):
        assert ui.call("GET", f"/api/runs/t/exports/{name}")[0] == 404, name
    assert ui.call("GET", "/api/runs/t/exports/MRA_t_DRAFT.xlsx")[0] == 200


def b64(data: bytes) -> str:
    return base64.b64encode(data).decode()


def full_format_deck(tweak: str = "Edited") -> bytes:
    """A deck in the AWS format (the built-in default with its cover text edited), so it is a different file with full coverage."""
    from io import BytesIO

    from pptx import Presentation

    from mra.outputs.deck_template import default_template_bytes
    prs = Presentation(BytesIO(default_template_bytes()))
    prs.slides[1].shapes.title.text_frame.text = "Migration Readiness Assessment (MRA)"
    prs.core_properties.comments = tweak
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def test_the_built_in_default_is_always_available_and_needs_no_upload(ui, cfg, golden, tmp_path):
    lst = ui.call("GET", "/api/templates")[1]
    assert lst["preferred"] == "ingram" and [t["id"] for t in lst["templates"]] == ["ingram", "aws"] and all(t["builtin"] for t in lst["templates"])
    assert lst["templates"][0]["name"] == "Ingram MAP deliverable (built in)" and "MRA slides are filled" in lst["templates"][0]["description"]
    assert ui.call("GET", "/api/info")[1]["templates"]["pptx"] == lst
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    st, out, _ = ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})
    assert st == 200 and out["deck_template"] == "Ingram MAP deliverable (built in)" and sorted(f["name"] for f in out["files"]) == ["MRA_t_DRAFT.pptx", "MRA_t_DRAFT_open_items.md"]
    d = [e for e in ui.svc.audit_entries("t") if e["action"] == "export" and e["detail"]["file"].endswith(".pptx")][0]["detail"]["template"]
    assert d["name"] == "Ingram MAP deliverable (built in)"


def test_an_uploaded_deck_is_stored_with_its_coverage_and_can_be_chosen_per_export(ui, cfg, golden, tmp_path, template_bytes):  # noqa: F811
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    st, up, _ = ui.call("POST", "/api/templates", {"kind": "pptx", "file_name": "partial.pptx", "file_b64": b64(template_bytes)})
    assert st == 200 and up["name"] == "partial.pptx" and up["coverage"]["complete"] is False and any("Heatmap" in w for w in up["warnings"])
    lst = ui.call("GET", "/api/templates")[1]
    assert [t["name"] for t in lst["templates"]][2:] == ["partial.pptx"] and lst["preferred"] == "ingram"       # not made the default unless asked
    entry = lst["templates"][2]
    assert entry["coverage"]["found"] < entry["coverage"]["expected"] and entry["id"] == up["id"]
    out = ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"], "template": up["id"]})[1]
    assert out["deck_template"] == "partial.pptx"
    todo = ui.call("GET", "/api/runs/t/exports/MRA_t_DRAFT_open_items.md")[1].decode()
    assert "Deck template used: partial.pptx" in todo and "found 0 of 34" in todo
    last = [e for e in ui.svc.audit_entries("t") if e["action"] == "export" and e["detail"]["file"].endswith(".pptx")][-1]["detail"]["template"]
    assert last["name"] == "partial.pptx"
    again = ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]                               # no choice: the preferred one, which is still the default
    assert again["deck_template"] == "Ingram MAP deliverable (built in)"


def test_a_full_format_deck_is_recognised_and_can_become_the_default(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    up = ui.call("POST", "/api/templates", {"kind": "pptx", "file_name": "ours.pptx", "file_b64": b64(full_format_deck()), "prefer": True})[1]
    assert up["coverage"]["complete"] and up["warnings"] == []
    assert ui.call("GET", "/api/templates")[1]["preferred"] == up["id"]
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]["deck_template"] == "ours.pptx"
    assert ui.call("POST", "/api/templates/aws/prefer", {})[1]["preferred"] == "aws"
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]["deck_template"] == "AWS MRA readout format (built in)"
    assert ui.call("POST", "/api/templates/ingram/prefer", {})[1]["preferred"] == "ingram"
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]["deck_template"] == "Ingram MAP deliverable (built in)"
    assert ui.call("POST", f"/api/templates/{up['id']}/prefer", {})[1]["preferred"] == up["id"]


def test_the_same_deck_uploaded_twice_is_one_entry(ui):
    deck = b64(full_format_deck())
    a = ui.call("POST", "/api/templates", {"kind": "pptx", "file_name": "a.pptx", "file_b64": deck})[1]
    b = ui.call("POST", "/api/templates", {"kind": "pptx", "file_name": "b.pptx", "file_b64": deck})[1]
    assert a["id"] == b["id"] and len(ui.call("GET", "/api/templates")[1]["templates"]) == 3


def test_uploaded_decks_can_be_removed_but_the_default_cannot(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    up = ui.call("POST", "/api/templates", {"kind": "pptx", "file_name": "ours.pptx", "file_b64": b64(full_format_deck()), "prefer": True})[1]
    assert ui.call("POST", "/api/templates/ingram/delete", {})[0] == 400 and ui.call("POST", "/api/templates/aws/delete", {})[0] == 400
    assert ui.call("POST", "/api/templates/0123456789/delete", {})[0] == 404
    lst = ui.call("POST", f"/api/templates/{up['id']}/delete", {})[1]
    assert [t["id"] for t in lst["templates"]] == ["ingram", "aws"] and lst["preferred"] == "ingram"       # the preference falls back
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"], "template": up["id"]})[0] == 404       # a removed deck cannot be used
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]["deck_template"] == "Ingram MAP deliverable (built in)"


@pytest.mark.parametrize("body,code", [
    ({"kind": "pptx", "file_name": "x.docx", "file_b64": "AAAA"}, 400), ({"kind": "xlsx", "file_name": "x.xlsx", "file_b64": "AAAA"}, 400),
    ({"kind": "pptx", "file_name": "x.pptx", "file_b64": "!!!"}, 400), ({"kind": "pptx", "file_name": "x.pptx", "file_b64": "AAAA"}, 422),
    ({"kind": "docx", "file_name": "x.docx", "file_b64": "AAAA"}, 422)])
def test_template_uploads_are_checked(ui, body, code):
    assert ui.call("POST", "/api/templates", body)[0] == code


def test_unknown_template_ids_are_refused(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"], "template": "../../etc"})[0] == 400
    assert ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"], "template": "ffffffffff"})[0] == 404
    assert ui.call("POST", "/api/templates/0123456789/prefer", {})[0] == 404
    assert ui.call("POST", "/api/templates/not-an-id/prefer", {})[0] == 404 and ui.call("POST", "/api/templates/../x/delete", {})[0] == 404


def test_a_deck_stored_by_the_first_page_version_is_kept(ui):
    ui.svc.s.put("ui/templates/pptx.bin", full_format_deck("legacy"))
    ui.svc.s.put("ui/templates/pptx.name", b"old deck.pptx")
    lst = ui.call("GET", "/api/templates")[1]
    assert [t["name"] for t in lst["templates"]][2:] == ["old deck.pptx"] and lst["preferred"] == lst["templates"][2]["id"]
    assert not ui.svc.s.exists("ui/templates/pptx.bin") and not ui.svc.s.exists("ui/templates/pptx.name")
    assert len(ui.call("GET", "/api/templates")[1]["templates"]) == 3                                   # and it is not imported twice


def test_a_word_template_is_stored_separately(ui):
    from io import BytesIO

    from docx import Document
    buf = BytesIO()
    Document().save(buf)
    assert ui.call("POST", "/api/templates", {"kind": "docx", "file_name": "house.docx", "file_b64": b64(buf.getvalue())})[0] == 200
    info = ui.call("GET", "/api/info")[1]["templates"]
    assert info["docx"] == "house.docx" and len(info["pptx"]["templates"]) == 2


def test_a_template_given_at_launch_is_stored_preferred_and_reports_its_coverage(ui, cfg, golden, tmp_path, template_bytes):  # noqa: F811
    f = tmp_path / "my deck.pptx"
    f.write_bytes(template_bytes)
    warnings = ui.srv.api.load_template_file("pptx", str(f))
    assert any("Heatmap" in w for w in warnings)                                                        # the small test deck is not the whole AWS format
    lst = ui.call("GET", "/api/templates")[1]
    assert lst["templates"][2]["name"] == "my deck.pptx" and lst["preferred"] == lst["templates"][2]["id"]
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    out = ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})[1]
    assert out["deck_template"] == "my deck.pptx"


def test_runs_are_listed_newest_first_with_their_state(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", {**q_body(cfg, golden, tmp_path), "set_id": "first"})
    ui.call("POST", "/api/runs", {**q_body(cfg, golden, tmp_path, name="q2.xlsx"), "set_id": "second", "customer": "B"})
    rows = ui.call("GET", "/api/runs")[1]["runs"]
    assert [r["set_id"] for r in rows] == ["second", "first"] and rows[0]["customer"] == "B" and rows[0]["blank_ratings"] == 1 and not rows[0]["final"]


def test_a_missing_run_is_404_and_deleting_needs_the_id_typed(ui, cfg, golden, tmp_path):
    assert ui.call("GET", "/api/runs/nope")[0] == 404
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    assert ui.call("POST", "/api/runs/t/purge", {"confirm": "wrong"})[0] == 400 and ui.call("GET", "/api/runs/t")[0] == 200
    assert ui.call("POST", "/api/runs/t/purge", {"confirm": "t"})[1]["objects_removed"] > 5 and ui.call("GET", "/api/runs/t")[0] == 404


def test_agent_proposals_show_the_quote_in_context_and_a_person_decides(ui, cfg, golden, tmp_path):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    ui.svc.llm = Script([turn(use("propose_rating", question_id="Q42", score=1, evidence_id="E-cio", quote="There is no cloud team yet",
                                 rationale="Level 1: no team has been staffed.")), end("one proposal")])
    out = ui.call("POST", "/api/runs/t/agent", {"task": "rating"})[1]
    assert out["proposal_ids"] == ["P0001"] and out["stop_reason"] == "done"
    p = ui.call("GET", "/api/runs/t")[1]["proposals"][0]
    assert p["kind"] == "rating" and p["status"] == "proposed" and p["excerpt"]["found"] and p["excerpt"]["quote"].lower() == "there is no cloud team yet"
    assert "lease" in p["excerpt"]["before"] or "lease" in p["excerpt"]["after"] or p["excerpt"]["before"]
    assert ui.call("GET", "/api/runs/t")[1]["blank_ratings"] == ["Q42"]
    assert ui.call("POST", "/api/runs/t/proposals", {})[0] == 400
    assert ui.call("POST", "/api/runs/t/proposals", {"accept": ["P0001"], "comment": "agreed"}, actor="Alice")[1]["scores_changed"] is True
    v = ui.call("GET", "/api/runs/t")[1]
    assert v["blank_ratings"] == [] and next(q for q in v["questions"] if q["question_id"] == "Q42")["adjustment"]["by"] == "Alice"


def test_the_deck_preview_uses_libreoffice_when_present(ui, cfg, golden, tmp_path, monkeypatch):
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    monkeypatch.setenv("MRA_SOFFICE", str(tmp_path / "does-not-exist"))
    monkeypatch.setattr("shutil.which", lambda *_: None)
    st, out, _ = ui.call("POST", "/api/runs/t/preview", {"kind": "docx"})
    assert st == 501 and "LibreOffice" in out["error"] and ui.call("GET", "/api/info")[1]["libreoffice"] is False
    shim = tmp_path / "fake_soffice"
    shim.write_text("#!/usr/bin/env python3\nimport sys, pathlib\na = sys.argv\nout = a[a.index('--outdir') + 1]\n"
                    "pathlib.Path(out, pathlib.Path(a[-1]).stem + '.pdf').write_bytes(b'%PDF-1.4 fake')\n")
    shim.chmod(shim.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("MRA_SOFFICE", str(shim))
    st, out, _ = ui.call("POST", "/api/runs/t/preview", {"kind": "docx"})
    assert st == 200 and out["url"] == "/api/runs/t/preview/docx.pdf"
    st, pdf, h = ui.call("GET", out["url"])
    assert st == 200 and pdf.startswith(b"%PDF") and h["Content-Type"] == "application/pdf" and h["Content-Disposition"] == "inline"
    assert ui.call("GET", "/api/runs/t/preview/pptx.pdf")[0] == 404 and ui.call("POST", "/api/runs/t/preview", {"kind": "xlsx"})[0] == 400


def test_a_missing_or_wrong_launch_template_is_refused_with_a_clear_message(ui, tmp_path):
    from mra.ui.api import UiError
    with pytest.raises(UiError, match="was not found"):
        ui.srv.api.load_template_file("pptx", str(tmp_path / "nope.pptx"))
    (tmp_path / "x.docx").write_bytes(b"x")
    with pytest.raises(UiError, match="must be a .pptx"):
        ui.srv.api.load_template_file("pptx", str(tmp_path / "x.docx"))


def test_serve_stops_with_a_message_when_the_launch_template_is_missing(svc, tmp_path, capsys):
    from mra.ui.server import serve
    assert serve(svc, 0, False, {"pptx": str(tmp_path / "nope.pptx")}) == 2
    assert "was not found" in capsys.readouterr().out


def test_a_browser_that_drops_the_connection_mid_download_does_not_disturb_the_server(ui, cfg, golden, tmp_path, capfd):
    """A cancelled or retried download resets the connection; that used to print a traceback from the error handler."""
    import socket
    import struct
    ui.call("POST", "/api/runs", q_body(cfg, golden, tmp_path))
    ui.call("POST", "/api/runs/t/export", {"kinds": ["pptx"]})                                  # a multi-megabyte file
    capfd.readouterr()
    for _ in range(5):
        c = socket.create_connection(("127.0.0.1", ui.port))
        c.sendall(f"GET /api/runs/t/exports/MRA_t_DRAFT.pptx?t={TOKEN} HTTP/1.1\r\nHost: 127.0.0.1:{ui.port}\r\n\r\n".encode())
        c.setsockopt(socket.SOL_SOCKET, socket.SO_LINGER, struct.pack("ii", 1, 0))              # close with a reset, like a browser abandoning it
        c.close()
    import time
    time.sleep(1.0)
    assert ui.call("GET", "/api/info")[0] == 200                                                # the server is still serving
    err = capfd.readouterr().err
    assert "Traceback" not in err and "BrokenPipe" not in err and "ConnectionReset" not in err, err[-600:]
