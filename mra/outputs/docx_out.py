"""Word report in the structure of the AWS 'Readout and Recommendations' document.

Cover, overview, contents, radar charts, heatmap, scores, details (observations and recommended actions for every
readiness activity), consolidated recommended actions, then every response. Report details that are not known yet are
printed as 'TBD', never left as template placeholders. A group with a blank rating is shown as Incomplete.
If a template .docx is supplied, its styles, header and footer are kept and its body is replaced.
"""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .charts import bars_png, radar_png

FILL = {"red": "F4B6B6", "amber": "FFE699", "green": "C6E0B4", "incomplete": "D9D9D9", "head": "1F3864"}
INK = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x59, 0x59, 0x59)


def _style(doc, name: str, fallback: str = "Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles[fallback]


def _heading(doc, text: str, level: int):
    try:
        return doc.add_heading(text, level)
    except KeyError:
        p = doc.add_paragraph()
        _run(p, text, bold=True, size={1: 14, 2: 11}.get(level, 10), color=INK)
        return p


def _shade(cell, hex_fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")      # CLEAR, never SOLID: SOLID renders black in some viewers
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _clear_body(doc) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _tag(p, source: str) -> None:
    _run(p, f"  [{source}]", italic=True, size=8, color=GREY)


def _table(doc, header: list[str], rows: list[list], widths: list[float], fills: dict | None = None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = _style(doc, "Table Grid").name
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        _shade(c, FILL["head"])
        c.paragraphs[0].text = ""
        c.paragraphs[0].paragraph_format.keep_with_next = True
        _run(c.paragraphs[0], h, bold=True, size=9, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].paragraphs[0].text = ""
            _run(cells[i].paragraphs[0], "" if v is None else str(v), size=9)
            if fills and (ri, i) in fills:
                _shade(cells[i], FILL.get(fills[(ri, i)], FILL["incomplete"]))
    return t


def _score(g: dict) -> str:
    return "Incomplete" if g["status"] != "complete" else f"{g['score']:.1f}"


def chart_images(m: dict) -> dict:
    """The two radar charts and one bar chart per workstream, as PNG bytes (shared with the deck writer)."""
    guide = m["radar_guide"]
    acts = {a["id"]: a for a in m["activities"]}
    out = {"perspectives": radar_png("Perspectives", [p["name"] for p in m["perspectives"]], [p["radar"] for p in m["perspectives"]], guide),
           "workstreams": radar_png("Workstreams", [w["name"] for w in m["workstreams"]], [w["radar"] for w in m["workstreams"]], guide),
           "bars": {}}
    for w in m["workstreams"]:
        rows = [("AVERAGE", w["score"], w["colour"])] + [(acts[i]["name"], acts[i]["score"], acts[i]["colour"]) for i in w["activity_ids"]]
        out["bars"][w["id"]] = bars_png(w["name"], rows)
    return out


def render_docx(m: dict, template: bytes | None = None) -> bytes:
    doc = Document(BytesIO(template)) if template else Document()
    if template:
        _clear_body(doc)
    else:
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)
        for name in ("Heading 1", "Heading 2", "Title"):
            st = doc.styles[name]
            st.font.name, st.font.color.rgb = "Arial", INK
        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = f"{m['status']}: {m['customer']}. Scores show cloud readiness, not the on-premises estate."
        footer.runs[0].font.size = Pt(8)
    bullet = _style(doc, "List Bullet")
    d = m["details"]
    tbd = lambda k: d.get(k) or "TBD"  # noqa: E731

    doc.add_paragraph(f"Migration Readiness Assessment: {m['customer']}", style=_style(doc, "Title").name)
    p = doc.add_paragraph()
    _run(p, m["status"] + ". ", bold=True, color=RGBColor(0xC0, 0x00, 0x00) if m["status"] == "DRAFT" else INK)
    _run(p, m["status_note"].removeprefix("DRAFT. "))
    doc.add_paragraph(f"Facilitated by: {tbd('facilitator')}")
    doc.add_paragraph("Migration Readiness Assessment: Readout and Recommendations")
    _heading(doc, "Overview", 1)
    _table(doc, ["Item", "Detail"], [["Customer", m["customer"]], ["Assessment name", "Migration Readiness Assessment"],
                                     ["Internal contact", tbd("internal_contact")], ["Sponsor contact", tbd("sponsor_contact")],
                                     ["Workshop date", tbd("workshop_date")], ["Readout date", tbd("readout_date")],
                                     ["Reference", m["set_id"]], ["Generated", m["generated_at"]], ["Text engine", m["engine"]]], [1.8, 4.7])
    if m["blank_ratings"]:
        p = doc.add_paragraph()
        _run(p, f"Incomplete: {len(m['blank_ratings'])} rating(s) are blank ({', '.join(m['blank_ratings'])}). "
                "Groups that contain them show as Incomplete until they are rated.", bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    _heading(doc, "Contents", 1)
    for line in ("Radar charts", "Heatmap", "Scores", "Details", "Consolidated recommended actions", "Responses"):
        doc.add_paragraph(line, style=bullet.name)

    imgs = chart_images(m)
    _heading(doc, "Radar charts", 1)
    for key in ("perspectives", "workstreams"):
        doc.add_picture(BytesIO(imgs[key]), width=Inches(4.4))
    _heading(doc, "Heatmap", 1)
    doc.add_paragraph("Green: readiness maturity needs minimal assistance. Amber: needs attention to accelerate momentum. "
                      "Red: requires immediate assistance to unblock progress. Grey: incomplete.")
    for w in m["workstreams"]:
        doc.add_picture(BytesIO(imgs["bars"][w["id"]]), width=Inches(6.0))

    _heading(doc, "Scores", 1)
    rows, fills = [], {}
    for p_ in m["perspectives"]:
        fills[(len(rows), 1)] = p_["colour"]
        rows.append([p_["name"], _score(p_), p_["answered"], p_["total"]])
    _table(doc, ["Perspective", "Score (1 to 5)", "Rated", "Questions"], rows, [2.6, 1.5, 1.0, 1.0], fills)
    doc.add_paragraph()
    rows, fills = [], {}
    for w in m["workstreams"]:
        fills[(len(rows), 2)] = w["colour"]
        rows.append([w["name"], "AVERAGE", _score(w)])
        for a in (x for x in m["activities"] if x["workstream"] == w["id"]):
            fills[(len(rows), 2)] = a["colour"]
            rows.append(["", a["name"], _score(a)])
    _table(doc, ["Workstream", "Readiness activity", "Score (1 to 5)"], rows, [2.2, 3.2, 1.2], fills)

    _heading(doc, "Details", 1)
    for w in m["workstreams"]:
        _heading(doc, w["name"], 2)
        for a in (x for x in m["activities"] if x["workstream"] == w["id"]):
            p = doc.add_paragraph()
            _run(p, f"{a['name']}  ({_score(a)}, {a['phase']})", bold=True, color=INK)
            p = doc.add_paragraph()
            _run(p, "Summary Observations: ", bold=True)
            _run(p, a["observation"]["text"])
            _tag(p, a["observation"]["source"])
            p = doc.add_paragraph()
            _run(p, "Recommended Actions: ", bold=True)
            if not a["actions"]:
                _run(p, "No action required." if not a["needs_action"] else "No action has been drafted yet.")
            for x in a["actions"]:
                q = doc.add_paragraph(style=bullet.name)
                _run(q, x["text"])
                _tag(q, x["source"])
                _run(q, f"  Owner: {x['owner'] or 'to be assigned'}; due: {x['due_date'] or 'to be set'}; {x['status'].lower()}.", size=8, color=GREY)
    if m["rejected_actions"]:
        _heading(doc, "Actions rejected by the reviewer", 2)
        for a in m["rejected_actions"]:
            doc.add_paragraph(f"{a['name']}: {a['comment']}", style=bullet.name)

    _heading(doc, "Consolidated recommended actions", 1)
    if m["action_rows"]:
        _table(doc, ["Workstream", "Readiness activity", "Phase", "Action", "Owner", "Due", "Status"],
               [[r["activity"]["workstream_name"], r["activity"]["name"], r["activity"]["phase"], r["action"]["text"],
                 r["action"]["owner"] or "TBD", r["action"]["due_date"] or "TBD", r["action"]["status"]] for r in m["action_rows"]],
               [0.9, 1.0, 0.7, 2.4, 0.7, 0.6, 0.6])
    else:
        doc.add_paragraph("No activity needs an action.")

    _heading(doc, "Responses", 1)
    _heading(doc, "Ratings responses", 2)
    qby = {q["question_id"]: q for q in m["questions"]}
    for w in m["workstreams"]:
        _heading(doc, w["name"], 3)
        for a in (x for x in m["activities"] if x["workstream"] == w["id"]):
            p = doc.add_paragraph()
            _run(p, a["name"], bold=True)
            _table(doc, ["Question", "Response", "Comments"],
                   [[m["question_bank"][i].text, "" if qby[i]["score"] is None else qby[i]["score"],
                     qby[i]["level_text"] or "(blank: no rating yet)"] for i in a["question_ids"]], [3.4, 0.8, 2.4])
    _heading(doc, "Text responses", 2)
    given = [t for t in m["text_answers"] if t["module"] is None]
    _table(doc, ["#", "Question", "Response", "Comments"],
           [[t["number"], t["question"], t["answer"] if t["type"] == "yes_no" else "", t["answer"] if t["type"] == "text" else ""]
            for t in given], [0.4, 2.6, 0.7, 2.9])
    inc = [t for t in m["text_answers"] if t["module"] is not None and t["included"]]
    if inc:
        _heading(doc, "Optional question responses", 2)
        _table(doc, ["#", "Question", "Response", "Comments"],
               [[t["number"], t["question"], t["answer"] if t["type"] == "yes_no" else "", t["answer"] if t["type"] == "text" else ""]
                for t in inc], [0.4, 2.6, 0.7, 2.9])
    _heading(doc, "Method and limits", 1)
    for line in ["Scores and colours are calculated by code from the agreed ratings: the plain mean per group, as in the AWS tool. "
                 "A model never sets or changes a score.",
                 "Text tagged model or stub was checked against the facts. Text tagged fallback failed those checks and was replaced.",
                 m["scope_note"], f"Versions: {', '.join(f'{k} {v}' for k, v in m['versions'].items())}; config hash {m['config_hash']}."]:
        doc.add_paragraph(line, style=bullet.name)
    for n in m["placeholder_notes"]:
        p = doc.add_paragraph(n, style=bullet.name)
        p.runs[0].italic = True
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
