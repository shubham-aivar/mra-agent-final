"""The Ingram Micro MAP deliverable deck, filled with the MRA results.

The client's template is used as it is: the same 25 slides in the same order, the same layouts, colours, fonts and pictures. Only these things change:
  * the MRA slides are filled (readiness charts, key recommendations, executive summary, areas of focus);
  * the sections that need other inputs (discovery, 7R, cost) are left blank: their slides stay, with their labels, but the sample data and the
    "<>" fill-in markers are removed, so no sample customer's numbers remain;
  * the red "Example" markers over the sample data are removed.
Nothing is added and no slide is removed. Every text we write is fitted to its box (see textfit); anything that had to be shortened is reported.
"""
from __future__ import annotations

import copy
import io
import re
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Pt

from ..focus import strongest, weakest
from . import textfit
from .pptx_out import _remove, _set_text, _update_charts

TEMPLATE_DIR = Path(__file__).resolve().parent / "templates"
INGRAM_NAME = "ingram_map_deliverable.pptx"
INGRAM_LABEL = "Ingram MAP deliverable (built in)"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
TITLES = {"title": "MAP Assessment", "exec": "Executive Summary", "scope": "Assessment Scope", "journey": "MAP Journey based on 7R's (Migrate)",
          "details": "7 R Analysis Details", "costs": "AWS Costs & Migration Savings Details", "recs": "MRA Readiness -Recommendations",
          "mra_exec": "MRA Executive Summary", "focus": "MRA Areas of Focus", "scores": "MRA Readiness Scores",
          "onprem": "Stay On Premise Cost Details"}
BLANK = ("exec", "scope", "journey", "details", "costs", "onprem")             # sections that need discovery, 7R or cost inputs
BLANK_WHY = {"exec": "discovery data (servers, storage, utilisation)", "scope": "discovery data (datacenters, cores, workloads, tools)",
             "journey": "the 7R analysis", "details": "the 7R analysis", "costs": "the AWS cost analysis",
             "onprem": "the on-premises cost analysis"}
LABELS = {"slides": "Slides kept", "anchors": "Sections recognised", "mra_charts": "MRA charts", "recommendations": "Key recommendations box",
          "exec_boxes": "Executive-summary boxes", "focus_boxes": "Areas-of-focus sentences"}


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s.replace("\u2019", "'").replace("\u2018", "'")).strip().lower()


def normalise(data: bytes) -> bytes:
    """A PowerPoint template (.potx) differs from a deck only in one content type: make it open as a deck."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        ct = z.read("[Content_Types].xml")
        if b"presentationml.template.main+xml" not in ct:
            return data
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in z.infolist():
                d = z.read(it.filename)
                zo.writestr(it, d.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                            if it.filename == "[Content_Types].xml" else d)
        return out.getvalue()


def ingram_template_bytes() -> bytes:
    return (TEMPLATE_DIR / INGRAM_NAME).read_bytes()


def _walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from _walk(sh.shapes)


def _first_lines(slide) -> list[tuple[str, float]]:
    """(first line of each text shape, its top in points): a title may share its shape with a subtitle."""
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame and sh.top is not None and sh.text_frame.text.strip():
            first = re.split(r"[\n\v]", sh.text_frame.text.strip())[0]
            if len(first) < 90:
                out.append((_norm(first), sh.top / 12700))
    return out


def locate(prs) -> dict[str, list]:
    found: dict[str, list] = {k: [] for k in TITLES}
    for slide in prs.slides:
        lines = _first_lines(slide)
        for key, title in TITLES.items():
            if any(t == _norm(title) and (key == "title" or top < 110) for t, top in lines):
                found[key].append(slide)
    return found


def is_ingram(data: bytes) -> bool:
    try:
        found = locate(Presentation(io.BytesIO(normalise(data))))
    except Exception:
        return False
    return sum(1 for k in ("title", "scope", "recs", "mra_exec", "scores", "journey") if found[k]) >= 4


def safe_right(slide, width_pt: float, margin: float = 6.0) -> float:
    """The right-hand limit for text on a slide: the left edge of the full-height decorative stripe the layout draws at the right edge."""
    limit = width_pt - 24.0
    for src in (slide.slide_layout, slide.slide_layout.slide_master):
        for sh in src.shapes:
            if sh.left is not None and not sh.is_placeholder and sh.left / 12700 > 0.8 * width_pt and sh.width / 12700 < 0.15 * width_pt \
                    and sh.height / 12700 > 360:
                limit = min(limit, sh.left / 12700 - margin)
    return limit


# ---- what to write ---------------------------------------------------------------------------------------------------------
def key_recommendations(m: dict, n: int = 5) -> list[str]:
    """The AWS standard actions behind the weakest readiness activities, one line each, weakest first; an action a reviewer rejected is not used."""
    rows = [(a["score"], i, a) for i, a in enumerate(m["activities"])
            if a["status"] == "complete" and a["needs_action"] and a["decision"] != "reject"]
    out: list[str] = []
    for _, _, a in sorted(rows, key=lambda t: (t[0], t[1])):
        for act in a["actions"]:
            line = (act.get("objective") or f"Improve {a['name']}").strip().rstrip(".") + "."
            if line not in out:
                out.append(line)
    return out[:n]


def strengths_and_gaps(m: dict, n: int = 4) -> tuple[list[str], list[str]]:
    pre = [a for a in m["activities"] if a["phase"] == "Pre-Mobilize"]
    return [a["name"] for a in strongest(pre, n)], [a["name"] for a in weakest(pre, n)]


def _answer(m: dict, number: int) -> list[str]:
    row = next((t for t in m["text_answers"] if t["number"] == number and t["included"]), None)
    return [ln.strip() for ln in (row["answer"] if row else "").splitlines() if ln.strip()]


# ---- filling -----------------------------------------------------------------------------------------------------------------
class Filler:
    def __init__(self, prs, m: dict):
        self.prs, self.m, self.notes, self.removed = prs, m, [], []
        self.found = locate(prs)
        self.stats = {"slides": len(prs.slides), "anchors": sum(1 for v in self.found.values() if v), "mra_charts": 0, "recommendations": 0,
                      "exec_boxes": 0, "focus_boxes": 0}

    @staticmethod
    def no_proofing(shape) -> None:
        """Our text carries no spelling or grammar marks (a squiggle under a word is mess in a client deliverable)."""
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                rpr = r._r.get_or_add_rPr()
                rpr.attrib.pop("err", None)
                rpr.set("noProof", "1")

    def put(self, shape, lines: list[str], label: str, collapse: bool = False, min_size: float = 10.0) -> None:
        """Write `lines` into `shape` keeping the template's formatting, fitted to the box."""
        self.clamp(shape)
        size = textfit.first_size(shape)
        sample = [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
        res = textfit.fit(shape, lines, sample, size, collapse=collapse, min_size=min_size)
        _set_text(shape, res["paragraphs"] or [""])
        textfit.apply_size(shape, res["size"], size)
        self.no_proofing(shape)
        self.m.setdefault("deck_fit", []).append({"box": label, "needed": res["needed"], "budget": res["budget"], "size": res["size"]})
        if res["shortened"]:
            self.notes.append(f"{label}: the text was shortened to fit its box on the slide (full text is in the Excel workbook)")

    def clamp(self, shape) -> None:
        """A box that reaches into the layout's right-hand stripe is narrowed to end before it, so its text wraps there instead of running into it."""
        slide = next(sl for sl in self.prs.slides if shape._element in sl.shapes._spTree.iter())
        limit = safe_right(slide, self.prs.slide_width / 12700) * 12700
        if shape.left + shape.width > limit:
            shape.width = int(limit - shape.left)

    def clear(self, shape) -> None:
        _set_text(shape, [""])

    def run(self) -> None:
        self.stats["mra_charts"] = _update_charts(self.prs, self.m)
        self.chart_labels()
        self.recommendations()
        self.exec_summary()
        self.focus()
        self.blank_sections()
        self.examples()

    # ---- slide: MRA Readiness - Recommendations -------------------------------------------------------------------------------
    def recommendations(self) -> None:
        for slide in self.found["recs"]:
            shape = next((sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip().startswith("Key Recommendations")), None)
            if shape is None:
                continue
            self.clamp(shape)
            paras = shape.text_frame.paragraphs
            bullets = [p for p in paras[2:] if p.text.strip()]
            if not bullets:
                continue
            body, proto = shape.text_frame._txBody, bullets[0]._p
            spacer = copy.deepcopy(next((p._p for p in paras[2:] if not p.text.strip()), proto))
            for r in spacer.findall(A + "r"):
                spacer.remove(r)
            sample = [p.text for p in bullets]
            size = textfit.first_size(shape, 16.0)
            for p in bullets[0]._p.iter(A + "rPr"):
                if p.get("sz"):
                    size = int(p.get("sz")) / 100
                    break
            recs = key_recommendations(self.m)
            res = textfit.fit(shape, ["Key Recommendations"] + recs, sample, size, indent_pt=22.5, gaps=max(0, len(recs) - 1) + 1, min_size=12.0)
            for p in paras[2:]:
                body.remove(p._p)
            for i, line in enumerate(res["paragraphs"][1:]):
                if i:
                    body.append(copy.deepcopy(spacer))
                node = copy.deepcopy(proto)
                for r in node.findall(A + "r")[1:]:
                    node.remove(r)
                run = node.find(A + "r")
                run.find(A + "t").text = line
                run.find(A + "rPr").attrib.pop("err", None)
                run.find(A + "rPr").set("noProof", "1")
                if res["size"] < size - 0.01:
                    run.find(A + "rPr").set("sz", str(int(res["size"] * 100)))
                body.append(node)
            self.stats["recommendations"] += 1
            fit = {"box": "Key recommendations", "needed": res["needed"], "budget": res["budget"], "size": res["size"]}
            self.m.setdefault("deck_fit", []).append(fit)
            if res["shortened"]:
                self.notes.append("Key recommendations: a line was shortened to fit its box on the slide")

    # ---- slide: MRA Executive Summary ------------------------------------------------------------------------------------------------
    def exec_summary(self) -> None:
        for slide in self.found["mra_exec"]:
            notes = sorted((sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "Notes"), key=lambda s: s.left)
            for shape, num, label in zip(notes, (2, 1), ("Strategy and Desired Outcomes (Q2)", "MRA Objectives (Q1)")):
                self.put(shape, _answer(self.m, num) or [""], label, collapse=True, min_size=11.0)       # readable first, the rest is in the workbook
                self.stats["exec_boxes"] += 1
            half = self.prs.slide_width / 2
            items = [sh for sh in slide.shapes if sh.has_text_frame and abs(sh.width / 12700 - 253) < 8 and abs(sh.height / 12700 - 23) < 8]
            strengths, gaps = strengths_and_gaps(self.m)
            for names, label, column in ((strengths, "Strength", [s for s in items if s.left < half]),
                                         (gaps, "Gap", [s for s in items if s.left >= half])):
                for i, shape in enumerate(sorted(column, key=lambda s: s.top)):
                    self.put(shape, [names[i]] if i < len(names) else [""], f"{label} {i + 1}")
                    self.stats["exec_boxes"] += 1

    # ---- slides: MRA Areas of Focus (perspectives, then workstreams) -------------------------------------------------------------------
    def focus(self) -> None:
        f = self.m["focus"]
        pairs = zip(self.found["focus"], (f["perspective_sentences"], f["workstream_sentences"]), ("Perspectives", "Workstreams"))
        for slide, sentences, label in pairs:
            boxes = sorted((sh for sh in slide.shapes if sh.has_text_frame and sh.name.startswith("Rectangle") and len(sh.text_frame.text) > 45),
                           key=lambda s: s.top)
            triangles = [sh for sh in slide.shapes if sh.name.startswith("Isosceles Triangle")]
            for i, shape in enumerate(boxes[:3]):
                self.put(shape, [sentences[i]] if i < len(sentences) else [""], f"Areas of focus ({label}) sentence {i + 1}", min_size=12.0)
                self.stats["focus_boxes"] += 1
                if i >= len(sentences) and triangles:                                      # a bullet marker with nothing beside it would look messy
                    mid = shape.top + shape.height / 2
                    marker = min(triangles, key=lambda t: abs(t.top + t.height / 2 - mid))
                    triangles.remove(marker)
                    _remove(marker)
                    self.removed.append(marker.name)

    # ---- the sections that need other inputs: blank, not deleted ----------------------------------------------------------------------------
    def blank_sections(self) -> None:
        self.m["deck_blank_sections"] = []
        for key in BLANK:
            for slide in self.found[key]:
                number = list(self.prs.slides).index(slide) + 1
                self.m["deck_blank_sections"].append(f"Slide {number} \"{TITLES[key]}\": left blank, it needs {BLANK_WHY[key]}")
                for sh in _walk(slide.shapes):
                    if sh.has_text_frame:
                        is_placeholder = "<>" in sh.text_frame.text                      # decided before any run is edited
                        for p in sh.text_frame.paragraphs:
                            for r in p.runs:
                                t = r.text
                                if is_placeholder:                                  # a value placeholder: clear the whole shape, units and all
                                    t = ""
                                if key == "scope" and t.strip() in ("Dates", "Tools used"):
                                    t = ""
                                if key == "details" and t.strip().startswith("Include details on discovered servers"):
                                    t = ""
                                if t != r.text:
                                    r.text = t
                    if getattr(sh, "has_chart", False) and sh.has_chart:
                        self.blank_chart(sh.chart)

    def chart_labels(self) -> None:
        """The template's charts carry hand-typed label text (a sample customer's "3.8"): remove it so labels show real values, one decimal."""
        ns = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
        for slide in self.prs.slides:
            for sh in _walk(slide.shapes):
                if not (getattr(sh, "has_chart", False) and sh.has_chart):
                    continue
                root = sh.chart._chartSpace
                if root.find(".//" + ns + "radarChart") is not None:
                    n = len(list(sh.chart.plots[0].categories))
                    sh.chart.category_axis.tick_labels.font.size = Pt(12 if n <= 6 else 10)
                for lbl in root.iter(ns + "dLbl"):
                    tx = lbl.find(ns + "tx")
                    if tx is not None:
                        lbl.remove(tx)
                if root.find(".//" + ns + "barChart") is None:
                    continue
                for dl in root.iter(ns + "dLbls"):
                    fmt = dl.find(ns + "numFmt")
                    if fmt is None:
                        fmt = dl.makeelement(ns + "numFmt", {})
                        anchor = next((c for c in dl if c.tag in (ns + "spPr", ns + "txPr", ns + "dLblPos", ns + "showLegendKey")), None)
                        (anchor.addprevious(fmt) if anchor is not None else dl.append(fmt))
                    fmt.set("formatCode", "0.0")
                    fmt.set("sourceLinked", "0")

    @staticmethod
    def blank_chart(chart) -> None:
        plot = chart.plots[0]
        data = CategoryChartData()
        data.categories = list(plot.categories)
        for s in plot.series:
            data.add_series(s.name, [None] * len(list(plot.categories)))
        chart.replace_data(data)
        if chart.has_title and chart.chart_title.has_text_frame:
            for p in chart.chart_title.text_frame.paragraphs:
                for r in p.runs:
                    r.text = textfit.clean_spaces(r.text.replace("N=<>", "").replace("<>", ""))
        chart.has_legend = False                                                    # an empty chart with a legend for nothing is mess

    def examples(self) -> None:
        for slide in self.prs.slides:
            for sh in list(slide.shapes):
                if sh.has_text_frame and sh.text_frame.text.strip() == "Example" and textfit.first_size(sh) >= 36:
                    self.removed.append(f"Example label on slide {list(self.prs.slides).index(slide) + 1}")
                    _remove(sh)


def structure_problems(template: Presentation, out: Presentation, allowed_removed: int) -> list[str]:
    """Nothing less, nothing more: the same slides, order and layouts, and the same shapes (minus the ones we remove on purpose)."""
    if len(template.slides) != len(out.slides):
        return [f"slide count {len(out.slides)} instead of {len(template.slides)}"]
    problems, removed = [], 0
    for i, (a, b) in enumerate(zip(template.slides, out.slides), 1):
        if a.slide_layout.name != b.slide_layout.name:
            problems.append(f"slide {i}: layout changed")
        na, nb = [s.name for s in a.shapes], [s.name for s in b.shapes]
        removed += len(na) - len(nb)
        if len(nb) > len(na) or any(n not in na for n in nb):
            problems.append(f"slide {i}: shapes were added")
    if removed > allowed_removed:
        problems.append(f"{removed} shapes removed, expected at most {allowed_removed}")
    return problems


def render_ingram(m: dict, template: bytes | None = None) -> bytes:
    data = normalise(template if template is not None else ingram_template_bytes())
    prs, base = Presentation(io.BytesIO(data)), Presentation(io.BytesIO(data))
    filler = Filler(prs, m)
    filler.run()
    problems = structure_problems(base, prs, len(filler.removed))
    if problems:
        raise ValueError("the deck would not match the template: " + "; ".join(problems))
    m["deck_stats"] = filler.stats
    m["deck_warnings"] = filler.notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def ingram_coverage(data: bytes, expected_data: bytes, probe) -> dict:
    """What the agent can fill in an Ingram-format deck, compared with the built-in one (counts of recognised sections, charts and text boxes)."""
    def stats(d: bytes) -> dict:
        m = probe()
        render_ingram(m, d)
        return m["deck_stats"]
    found, expected = stats(data), stats(expected_data)
    items, warnings = [], []
    for key, label in LABELS.items():
        f, e = int(found.get(key, 0)), int(expected.get(key, 0))
        items.append({"key": key, "label": label, "found": f, "expected": e, "ok": f >= e})
        if f < e:
            warnings.append(f"{label}: found {f} of {e}; the rest will not be filled")
    return {"items": items, "found": sum(i["found"] for i in items), "expected": sum(i["expected"] for i in items),
            "complete": all(i["ok"] for i in items), "warnings": warnings, "format": "ingram"}
